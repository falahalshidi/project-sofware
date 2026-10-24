from pathlib import Path
from shutil import copy2
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import math

BASE=Path('/Users/falahalshidi/Downloads/dev/project sofware ')
OUT=BASE/'Software Construction Report'
D=OUT/'Project Diagrams'; S=OUT/'Supporting Files'
D.mkdir(parents=True,exist_ok=True); S.mkdir(parents=True,exist_ok=True)
REG='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'; BOLD='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
if not Path(REG).exists(): REG='/System/Library/Fonts/Supplemental/Arial.ttf'; BOLD='/System/Library/Fonts/Supplemental/Arial Bold.ttf'
def ft(n,b=False): return ImageFont.truetype(BOLD if b else REG,n)
NAVY='#273746'; TEAL='#55AEB7'; PALE='#EEF6F6'; GOLD='#FFF8E4'; GREEN='#EEF7F0'; RED='#FDE8E8'; WHITE='#FFFFFF'; DARK='#1E2932'; MUTED='#69727A'
def make(title):
    im=Image.new('RGB',(1800,1000),WHITE); dr=ImageDraw.Draw(im); dr.rectangle((0,0,1800,86),fill=NAVY); dr.text((55,25),title,font=ft(34,True),fill=WHITE); return im,dr
def centre(dr,xy,text,font=ft(24,True),fill=DARK):
    x1,y1,x2,y2=xy; bb=dr.multiline_textbbox((0,0),text,font=font,spacing=5,align='center'); dr.multiline_text(((x1+x2-bb[2]+bb[0])/2,(y1+y2-bb[3]+bb[1])/2),text,font=font,fill=fill,align='center',spacing=5)
def bx(dr,xy,text,fill=PALE,outline=TEAL,font=ft(24,True)):
    dr.rounded_rectangle(xy,20,fill=fill,outline=outline,width=4); centre(dr,xy,text,font)
def ar(dr,a,b,label=''):
    dr.line([a,b],fill=NAVY,width=5); ang=math.atan2(b[1]-a[1],b[0]-a[0]); p1=(b[0]-22*math.cos(ang-.45),b[1]-22*math.sin(ang-.45)); p2=(b[0]-22*math.cos(ang+.45),b[1]-22*math.sin(ang+.45)); dr.polygon([b,p1,p2],fill=NAVY)
    if label:
        x=(a[0]+b[0])//2; y=(a[1]+b[1])//2; dr.rounded_rectangle((x-90,y-22,x+90,y+22),10,fill=WHITE); centre(dr,(x-90,y-22,x+90,y+22),label,ft(18),MUTED)
def diagram(name,title,fn):
    im,dr=make(title); fn(dr); path=D/(name+'.png'); im.save(path,dpi=(180,180)); return path
def arch(dr):
    bx(dr,(100,250,480,570),'Visitor\nBrowser\nReact UI'); bx(dr,(690,250,1090,570),'Node / Express\nAPI Server'); bx(dr,(1300,130,1700,300),'Open-Meteo\nWeather + Search',GOLD,'#D6B84C'); bx(dr,(1300,390,1700,560),'Mibyan AI\nChat Completions',GOLD,'#D6B84C'); bx(dr,(1300,650,1700,820),'Nominatim +\nOpenStreetMap',GREEN,'#70A77A'); ar(dr,(480,410),(690,410),'HTTP JSON'); ar(dr,(1090,330),(1300,215),'weather'); ar(dr,(1090,450),(1300,475),'chat'); ar(dr,(1090,530),(1300,735),'map / GPS')
def usecase(dr):
    dr.ellipse((70,370,260,560),fill=RED,outline='#B45D5D',width=4); centre(dr,(70,370,260,560),'Visitor',ft(25,True)); dr.rounded_rectangle((380,140,1510,840),25,outline=NAVY,width=5,fill='#FCFCFB'); dr.text((945,165),'SkyCast Weather System',font=ft(30,True),fill=NAVY,anchor='mm')
    for i,t in enumerate(['Search for a city','View current weather','Browse hourly forecast','Manage saved places','Use current location','Ask the weather assistant']):
        y=250+i*90; dr.ellipse((610,y-31,1280,y+31),fill=PALE,outline=TEAL,width=3); centre(dr,(610,y-31,1280,y+31),t,ft(22)); ar(dr,(260,y),(610,y),'uses')
def journey(dr):
    xs=[80,370,660,950,1240,1530]
    vals=[('Entry','Open site\nNo login'),('Locate','GPS or\nsearch'),('Explore','Current +\nhourly'),('Decide','Activities +\nmap'),('Ask','Weather\nassistant'),('Return','Saved\nplaces')]
    for i,(h,b) in enumerate(vals):
        x=xs[i]; bx(dr,(x,300,x+210,650),h+'\n\n'+b,PALE if i%2==0 else GOLD,TEAL if i%2==0 else '#D6B84C',ft(21,True))
        if i<5: ar(dr,(x+210,475),(x+280,475))
def flow(dr):
    bx(dr,(100,350,420,500),'Enter'); bx(dr,(570,350,930,500),'Choose location'); bx(dr,(1090,205,1470,335),'Search city'); bx(dr,(1090,520,1470,650),'Use GPS'); bx(dr,(1540,350,1740,500),'Load dashboard'); ar(dr,(420,425),(570,425),'start'); ar(dr,(930,390),(1090,270),'search'); ar(dr,(930,460),(1090,585),'GPS'); ar(dr,(1470,270),(1540,425),'place'); ar(dr,(1470,585),(1540,425),'place')
def activity(dr):
    bx(dr,(100,350,450,500),'Start',RED,'#B45D5D'); bx(dr,(650,350,1120,500),'Read weather'); bx(dr,(1320,180,1710,320),'Risk high?',GOLD,'#D6B84C'); bx(dr,(1320,500,1710,640),'Indoor suggestion',RED,'#B45D5D'); bx(dr,(1320,800,1710,940),'Outdoor suggestion',GREEN,'#70A77A'); ar(dr,(450,425),(650,425)); ar(dr,(1120,425),(1320,250),'evaluate'); ar(dr,(1515,320),(1515,500),'yes'); ar(dr,(1515,640),(1515,800),'no')
def sequence(dr):
    xs=[170,520,900,1270,1600]; names=['Visitor','React UI','Express API','Open-Meteo','Mibyan']
    for x,n in zip(xs,names): dr.text((x,145),n,font=ft(23,True),fill=NAVY,anchor='mm'); dr.line((x,180,x,910),fill='#B8C3C7',width=3)
    for a,b,l,y in [((170,250),(520,250),'search',250),((520,330),(900,330),'GET /search',330),((900,410),(1270,410),'geocode',410),((1270,490),(900,490),'places',490),((170,650),(520,650),'ask',650),((520,730),(900,730),'POST /chat',730),((900,810),(1600,810),'completion',810)]: ar(dr,a,b,l)
def component(dr):
    data=[('App.tsx','UI + state',100,230),('styles.css','responsive styles',100,520),('server/index.ts','API routes',700,230),('types.ts','data contracts',700,520),('localStorage','anonymous state',1300,230),('External APIs','weather / map / AI',1300,520)]
    for n,t,x,y in data: bx(dr,(x,y,x+390,y+180),n+'\n\n'+t,PALE if x<1000 else GOLD,TEAL if x<1000 else '#D6B84C',ft(21,True))
    ar(dr,(490,320),(700,320),'imports'); ar(dr,(490,610),(700,610),'types'); ar(dr,(1090,320),(1300,320),'state'); ar(dr,(1090,610),(1300,610),'HTTP')
def cycle(dr):
    pts=[(900,180),(1450,380),(1250,780),(550,780),(350,380)]; labs=['Requirements','Design','Implementation','Verification','Evaluation']
    for p,l in zip(pts,labs): dr.ellipse((p[0]-145,p[1]-65,p[0]+145,p[1]+65),fill=PALE,outline=TEAL,width=4); centre(dr,(p[0]-145,p[1]-65,p[0]+145,p[1]+65),l,ft(22,True))
    for a,b in zip(pts,pts[1:]+pts[:1]): ar(dr,a,b)
    dr.text((900,500),'Iterative construction',font=ft(32,True),fill=NAVY,anchor='mm')
paths=[
diagram('01_system_architecture','Figure 1 System Architecture Diagram',arch),
diagram('02_use_case','Figure 2 Use Case Diagram',usecase),
diagram('03_customer_journey','Figure 3 Proposed Customer Journey Map',journey),
diagram('04_user_flow','Figure 4 User Flow Diagram',flow),
diagram('05_activity','Figure 5 Weather Activity Diagram',activity),
diagram('06_sequence','Figure 6 Search and Chat Sequence Diagram',sequence),
diagram('07_component','Figure 7 Component Diagram',component),
diagram('08_development_cycle','Figure 8 Iterative Development Cycle',cycle)]
if (BASE/'report-mobile.png').exists(): copy2(BASE/'report-mobile.png',S/'mobile-home-and-chat.png')

def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),color)
def borders(cell,color='D9D9D9'):
    tcPr=cell._tc.get_or_add_tcPr(); bs=tcPr.first_child_found_in('w:tcBorders')
    if bs is None: bs=OxmlElement('w:tcBorders'); tcPr.append(bs)
    for e in ['top','left','bottom','right','insideH','insideV']:
        el=bs.find(qn('w:'+e))
        if el is None: el=OxmlElement('w:'+e); bs.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:color'),color)
def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; borders(cell)
            if ri==0: shade(cell,'273746')
            elif ri%2==0: shade(cell,'F3F7F7')
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                for r in p.runs:
                    r.font.name='Arial'; r.font.size=Pt(8.2); r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(35,43,49); r.font.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
def para(doc,text,style='Body Text'):
    p=doc.add_paragraph(style=style); p.add_run(text); return p
def bullets(doc,items):
    for x in items: doc.add_paragraph(x,style='List Bullet')
def fig(doc,path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.35)); c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(80,90,96)
def br(doc): doc.add_page_break()

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(35,43,49)
for nm,size in [('Title',28),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
    st=styles[nm]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_before=Pt(13); st.paragraph_format.space_after=Pt(6); st.paragraph_format.keep_with_next=True
styles['Body Text'].font.name='Arial'; styles['Body Text'].font.size=Pt(10.5); styles['Body Text'].paragraph_format.line_spacing=1.12; styles['Body Text'].paragraph_format.space_after=Pt(6)
sec.header.paragraphs[0].text='COMP3131  |  Software Construction  |  Al-Tarqi SkyCast'; sec.header.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.add_run('Sohar University  •  '); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(115); p.add_run('Al-Tarqi SkyCast Weather Web Application')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Software Construction Project Report').bold=True
for line in ['Sohar University','College of Computing and Information Technology','Course: Software Construction','Course Code: COMP3131','Student Name: [Student Name]','Student ID: [Student ID]','Instructor: [Instructor Name]','Academic Semester: [Semester]','Submission Date: [Submission Date]']:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4); p.add_run(line).font.size=Pt(11)
para(doc,'Prepared from the implemented React, TypeScript, Node.js and Express project in the shared workspace.').alignment=WD_ALIGN_PARAGRAPH.CENTER; br(doc)

doc.add_heading('Abstract',0)
para(doc,'This report documents Al-Tarqi SkyCast, a responsive weather web application presented through the SkyCast interface. The system addresses the practical need for a concise, location-aware weather overview and a simple path from city search to an outdoor decision. The implemented application opens without authentication, obtains current and forecast data from Open-Meteo, supports Arabic and English place search, resolves browser coordinates through reverse geocoding, displays current and hourly conditions, stores favourite places in browser localStorage, renders an interactive Leaflet map, and provides a weather assistant through a server-side Mibyan AI adapter. The server protects provider credentials and exposes application-specific routes to the React client.')
para(doc,'The project uses React, TypeScript, Vite, Node.js, Express, Leaflet, OpenStreetMap tiles, Lucide icons and CSS. The assistant uses Mibyan’s OpenAI-compatible Chat Completions contract with a weather-aware prompt and a deterministic fallback when the provider is unavailable or the project key is unauthorised. This distinction is important: the integration is implemented, but the current workspace key was rejected during live verification, so the report does not claim successful model output.')
para(doc,'The report separates verified implementation facts from future proposals. It covers requirements, architecture, UML-style diagrams, customer journey, UX, implementation, testing, quality, security, limitations and future work. It demonstrates Software Construction through clear API boundaries, environment configuration, reusable UI functions, error handling and iterative verification.')
doc.add_heading('Table of Contents',0)
table(doc,['Section','Page'],[['Abstract','2'],['1 Introduction','3'],['2 Problem Analysis','3'],['3 Objectives and Scope','4'],['4 Stakeholder Analysis','4'],['5 Requirements Analysis','4–5'],['6 Development Methodology','5–6'],['7 System Architecture','6–7'],['8 Use Case Analysis','7–8'],['9 Customer Journey and User Flow','8–9'],['10 UML and Structural Diagrams','9–10'],['11 User Interface and UX Design','11–12'],['12 Database Design','13'],['13 Technology Stack','13'],['14 Software Implementation','13–14'],['15 Software Construction Fundamentals','14'],['16 Testing Strategy','15'],['17 Software Quality Evaluation','15'],['18 Security and Privacy','16'],['19 Challenges and Limitations','16'],['20 Future Improvements','16'],['21 Conclusion','16'],['22 References','17'],['Appendices','17–18']]); br(doc)

doc.add_heading('1 Introduction',1)
doc.add_heading('1.1 Background',2); para(doc,'Al-Tarqi SkyCast is a browser-based weather dashboard whose visible product identity is SkyCast Weather. It was constructed as an interactive utility rather than a static HTML mock-up. The visitor enters directly, chooses a place, reads the current and near-term weather, saves useful places and asks the assistant for an interpretation.')
doc.add_heading('1.2 Problem Statement',2); para(doc,'Weather information is often fragmented between a current-temperature screen, a forecast view, a search form and a map. The project addresses this fragmentation by placing the most relevant information in one visual dashboard. The problem is framed as an engineering requirement: a visitor needs a fast, readable and location-aware way to interpret current weather and plan an activity.')
doc.add_heading('1.3 Aim and Objectives',2); para(doc,'The aim is to construct a responsive weather tool that gives visitors a reliable current view, supports place exploration and explains the available data in natural language when the AI provider is available.')
bullets(doc,['Open directly without registration.','Search places in Arabic and English with country and region disambiguation.','Use browser location only after an explicit user action.','Display current temperature, apparent temperature, humidity, high/low, sunrise/sunset and hourly data.','Persist favourite places and theme preference locally.','Display an interactive map and selected-place marker.','Connect the weather assistant through a server-side Mibyan adapter.','Keep the hourly strip scrollable inside the page on phone widths.'])
doc.add_heading('1.4 Scope and Report Structure',2); para(doc,'The implemented scope is a single-page dashboard and Express API. It excludes accounts, database records, payments, push notifications, administrative screens and production deployment. Later sections document the actual code and distinguish future work from delivered behaviour.')
doc.add_heading('2 Problem Analysis',1); para(doc,'The primary user is an anonymous visitor who wants an immediate answer. The interface therefore privileges a current card, a compact hourly timeline and activity suggestions. Search results include region and country because city names are not necessarily unique. Favourites are deliberately device-local because no identity or database layer exists.')
table(doc,['Need','Implemented response','Boundary'],[['Immediate orientation','Current card and selected city','No personal profile'],['Place discovery','Arabic/English geocoding','Provider quality varies'],['Comparison','Saved rows with temperatures','No advanced analytics'],['Interpretation','Activity cards and chat','Informational only'],['Spatial context','Leaflet map marker','No weather raster layers']])
doc.add_heading('3 Objectives and Scope',1)
table(doc,['ID','Objective','Success criterion','Status'],[['OBJ-01','Fast access','Root dashboard opens without login','Implemented'],['OBJ-02','Current weather','Values populated from /api/weather','Implemented'],['OBJ-03','Global search','Arabic/English results include region','Implemented'],['OBJ-04','Saved places','localStorage survives refresh','Implemented'],['OBJ-05','Map context','Marker updates after selection','Implemented'],['OBJ-06','Assistant','Valid Mibyan key returns completion','Partially verified'],['OBJ-07','Accounts and alerts','Future backend scope','Future']])
para(doc,'“Partially verified” means the adapter and request path exist, but live Mibyan output was not confirmed because the current project key was rejected. This is not reported as a successful AI result.')
doc.add_heading('4 Stakeholder Analysis',1)
table(doc,['Stakeholder','Needs','Impact'],[['Website visitor','Fast weather, search, readable UI, no account','High'],['Student developer','Maintainable code and provider configuration','High'],['Course instructor','Evidence of construction, testing and limitations','Medium'],['External providers','Valid requests and credential protection','Medium'],['Future maintainer','Documented endpoints and modular boundaries','Medium']])
para(doc,'No administrator, merchant or service-provider workflow was found in the inspected source, so these roles are not claimed as implemented stakeholders.')
doc.add_heading('5 Requirements Analysis',1)
doc.add_heading('5.1 Functional Requirements',2)
fr=[['FR-01','Display selected-place weather summary','Visitor','Current card and metrics render','Implemented','UI/API test'],['FR-02','Search a place in Arabic or English','Visitor','Disambiguated results appear','Implemented','Integration test'],['FR-03','Select result and reload weather','Visitor','Cards and map update','Implemented','System test'],['FR-04','Request device location','Visitor','Reverse-geocoded area is shown','Implemented; permission-dependent','Browser test'],['FR-05','Browse 24-hour strip internally','Visitor','Page does not scroll horizontally','Implemented','Responsive test'],['FR-06','Add/remove saved places','Visitor','localStorage list updates','Implemented','Functional test'],['FR-07','Toggle units and theme','Visitor','Visible state changes','Partially implemented','UI test'],['FR-08','Show map marker','Visitor','Marker follows place','Implemented','Integration test'],['FR-09','Ask assistant question','Visitor','Model or fallback answer appears','Partially verified','API test'],['FR-10','Handle provider failure safely','System','Fallback without key exposure','Implemented','Failure test']]
table(doc,['ID','Requirement','Actor','Expected outcome','Status','Verification'],fr)
doc.add_heading('5.2 Non-Functional Requirements',2)
table(doc,['Attribute','Interpretation','Evidence/status'],[['Usability','Understand dashboard without registration','Mobile/browser review; verified for tested view'],['Performance','Keep client work and provider routes concise','No benchmark; not fully tested'],['Reliability','Maintain useful chat when provider fails','Fallback observed'],['Maintainability','Separate client, server, types and styles','Verified by source inspection'],['Compatibility','Work at desktop and phone widths','Phone viewport checked'],['Security','Keep keys server-side','Mibyan key read by Express only'],['Accessibility','Readable contrast and labelled controls','Partial; full audit not executed']])
doc.add_heading('5.3 Traceability Matrix',2)
table(doc,['Requirement','Component','Test case','Status'],[['FR-01','App.tsx current card','TC-01','Verified'],['FR-02','/api/search + ask bar','TC-02','Verified'],['FR-04','useLocation + reverse-geocode','TC-03','Permission-dependent'],['FR-05','hour-columns CSS','TC-04','Verified'],['FR-06','favorites/localStorage','TC-05','Verified'],['FR-08','WeatherMap','TC-06','Implemented'],['FR-09','/api/chat + Mibyan','TC-07','External key failed'],['FR-10','Server fallback','TC-08','Verified']])
doc.add_heading('6 Development Methodology',1)
para(doc,'No formal Scrum board, sprint record or signed lifecycle document was found. It would therefore be inaccurate to claim a formal Agile, Scrum or Waterfall process. The available evidence shows an iterative construction loop: visual feedback was converted into requirements, bounded code changes were made, builds and browser checks were executed, and external-provider failures were diagnosed.')
fig(doc,D/'08_development_cycle.png','Figure 8: Iterative construction cycle inferred from the project work')
doc.add_heading('7 System Architecture',1)
para(doc,'The browser contains the React presentation layer and calls local Express routes. Express calls Open-Meteo for weather and geocoding, Nominatim for reverse geocoding, and Mibyan for chat. Leaflet loads OpenStreetMap tiles from the browser. Provider keys remain on the server.')
fig(doc,D/'01_system_architecture.png','Figure 1: System architecture and external service boundaries')
table(doc,['Component','Responsibility','Evidence'],[['React/Vite client','UI, state, interactions','src/App.tsx and styles.css'],['Express server','Provider boundary and validation','server/index.ts'],['localStorage','Anonymous place/theme persistence','skycast:* keys'],['Open-Meteo','Forecast and geocoding','/api/weather and /api/search'],['Nominatim','Coordinate-to-area name','/api/reverse-geocode'],['Mibyan AI','Weather chat completion','/api/chat and MIBYAN_*'],['Leaflet/OpenStreetMap','Map and tiles','WeatherMap']])
doc.add_heading('8 Use Case Analysis',1)
para(doc,'The primary actor is Visitor. The visitor is anonymous; there is no implemented account or administrator actor.')
fig(doc,D/'02_use_case.png','Figure 2: Use case diagram for the Al-Tarqi SkyCast website')
for h,t in [('UC-01 Search for a place','The visitor enters Arabic or English text; the client debounces it; Express queries both language variants; results are deduplicated by coordinates; the visitor selects a result. Empty results and provider failure are handled without crashing the page.'),('UC-02 View current weather','After a place is selected, the client loads current, hourly and daily payloads and renders the cards. Provider failure resolves loading without inventing values.'),('UC-03 Use current location','The visitor presses the location action, grants permission, and receives reverse-geocoded coordinates. Permission denial or reverse-geocoding failure produces a safe generic label.'),('UC-04 Ask the assistant','The visitor submits a question. The server sends location and weather context to Mibyan. If the provider key is missing or rejected, a bounded weather-data fallback is returned.'),('UC-05 Manage saved places','The visitor adds a search result or removes a saved row. React state and localStorage are updated, with duplicate place IDs ignored.')]:
    doc.add_heading(h,3); para(doc,t)
doc.add_heading('9 Customer Journey and User Flow',1)
para(doc,'No field research or analytics dataset was supplied. The following journey is therefore proposed and illustrative, not a measured customer study.')
fig(doc,D/'03_customer_journey.png','Figure 3: Proposed customer journey map')
table(doc,['Stage','Action','Goal','Touchpoint','Pain point','Opportunity'],[['Entry','Open site','Get weather now','Root dashboard','Location may be unclear','Show selected city clearly'],['Locate','Use GPS/search','Find relevant place','Location/search control','Similar names','Show country/region'],['Explore','Read cards','Understand now/next hours','Current and hourly cards','Information density','Preserve hierarchy'],['Decide','Review activities/map','Plan activity','Activity icons/map','Suggestions are informational','Label clearly'],['Ask','Send question','Interpret data','Chat popup','External provider limits','Show provider/fallback status'],['Return','Open saved place','Compare quickly','localStorage rows','Device-specific data','Optional account sync later']])
fig(doc,D/'04_user_flow.png','Figure 4: User flow from entry to selected weather dashboard')
doc.add_heading('10 UML and Structural Diagrams',1)
doc.add_heading('10.1 Activity Diagram',2); para(doc,'The activity model shows how current weather can inform a simple activity suggestion. It is informational and is not a safety or medical decision engine.'); fig(doc,D/'05_activity.png','Figure 5: Activity diagram for an activity suggestion')
doc.add_heading('10.2 Sequence Diagram',2); para(doc,'The sequence diagram shows search and chat interactions across the visitor, React UI, Express API and external providers.'); fig(doc,D/'06_sequence.png','Figure 6: Search and chat sequence diagram')
doc.add_heading('10.3 Component Diagram',2); para(doc,'Because the project is functional React/TypeScript rather than a class-based object model, a component diagram is more accurate than an invented class diagram.'); fig(doc,D/'07_component.png','Figure 7: Component diagram of the implemented code structure')
doc.add_heading('11 User Interface and UX Design',1)
para(doc,'The visual language is a light, card-based dashboard inspired by the supplied Atmos Weather reference. It uses a full-width canvas, rounded paper cards, restrained grey surfaces, dark ink typography and teal accents. The top bar combines branding, theme/unit controls, city search and current-location access. The main area prioritises today’s view, a five-card viewport over a horizontally scrollable 24-hour strip, the current-weather card, activities, map and saved places.')
para(doc,'At phone width the dashboard becomes a vertical flow. The hourly strip keeps its own horizontal scroll so the document itself does not expand. The chat becomes a bounded popup. Browser inspection at a 539px viewport showed document width equal to viewport width while the hourly strip retained a larger internal scroll width.')
para(doc,'Accessibility measures include readable contrast, semantic buttons, Escape-to-close search, aria labels on key chat controls and a design that does not require colour alone to understand weather values. A complete WCAG and screen-reader audit was not executed.')
if (S/'mobile-home-and-chat.png').exists(): fig(doc,S/'mobile-home-and-chat.png','Figure 9: Actual mobile browser screenshot of the dashboard and chat popup')
doc.add_heading('12 Database Design',1)
para(doc,'No database was found. There are no migrations, collections, server-side user records or authentication routes. Saved places, selected place and theme are stored in browser localStorage. This is suitable for anonymous local persistence but is device-specific and cannot synchronise between devices.')
doc.add_heading('12.1 Proposed Future Database',2); para(doc,'If accounts are later required, a future schema could contain User, SavedPlace, Preference and ChatSession entities with strict ownership checks. This is a proposal only and is not an implemented database.')
doc.add_heading('13 Technology Stack',1)
table(doc,['Technology','Purpose','Where used','Justification'],[['React','Interactive UI','src/App.tsx','Component rendering and hooks'],['TypeScript','Types and checks','src, server','Catches interface mistakes early'],['Vite','Dev/build tooling','package scripts','Fast feedback and bundling'],['Node/Express','Server API','server/index.ts','Protects keys and centralises calls'],['Open-Meteo','Weather/geocoding','API routes','Forecast payload'],['Mibyan AI','Chat','/api/chat','OpenAI-compatible endpoint'],['Leaflet/OSM','Map','WeatherMap','Interactive marker map'],['Lucide React','Icons','App.tsx','Reusable SVG icons'],['localStorage','Local persistence','browser effects','No DB needed currently']])
doc.add_heading('14 Software Implementation',1)
doc.add_heading('14.1 Structure',2); para(doc,'App.tsx contains the main composition, state and interaction handlers. styles.css contains the visual system and responsive overrides. types.ts defines Place and WeatherData. server/index.ts contains Express routes and external-provider calls. package.json provides client, server, combined development and production build scripts.')
table(doc,['File','Responsibility'],[['src/App.tsx','Dashboard, state, weather/search/chat interactions, Metric and WeatherMap'],['src/styles.css','Typography, cards, responsive layout, theme and chat'],['src/types.ts','Place and WeatherData contracts'],['server/index.ts','Health, search, weather, reverse-geocode and chat routes'],['.env/.env.example','Runtime configuration and key placeholders'],['package.json','Dependencies and scripts']])
doc.add_heading('14.2 State and Data',2); para(doc,'Hooks manage selected place, favourites, weather, search text, unit, theme, chat messages and typing state. Effects load weather after a place change, debounce search, load favourite temperatures, write localStorage and scroll the chat to the latest message.')
doc.add_heading('14.3 API Routes',2); table(doc,['Route','Method','Purpose','Dependency'],[['/api/health','GET','Health/config response','None'],['/api/search','GET','Arabic/English place search','Open-Meteo geocoding'],['/api/weather','GET','Current/hourly/daily forecast','Open-Meteo'],['/api/reverse-geocode','GET','Coordinate area name','Nominatim'],['/api/chat','POST','Weather assistant','Mibyan or fallback']])
doc.add_heading('14.4 Error Handling and Reuse',2); para(doc,'The client catches failed fetches and prevents a failed request from crashing the page. The server validates coordinates and uses a bounded chat fallback. Metric removes repeated metric markup, WeatherMap isolates map lifecycle, environment variables keep provider configuration out of browser code, and CSS variables reuse theme values. App.tsx is still relatively concentrated and should be decomposed as the feature set grows.')
doc.add_heading('15 Software Construction Fundamentals',1)
for h,t in [('15.1 Minimizing Complexity','The system avoids accounts and a database because the current anonymous workflow does not need them. Its API surface is small, although App.tsx is now the main complexity hotspot.'),('15.2 Anticipating Change','Provider URLs and model names are environment-configurable. The API boundary permits upstream replacement without rewriting every UI card.'),('15.3 Constructing for Verification','The build performs TypeScript and Vite checks; the health endpoint exposes a simple diagnostic; browser checks cover Arabic search, phone layout, internal hourly scrolling and chat behaviour.'),('15.4 Reuse','React, Express, Leaflet and Lucide are reused libraries. Metric, WeatherMap, CSS variables and typed data shapes reduce repetition.'),('15.5 Standards in Construction','The project uses TypeScript, npm scripts, a src/server separation and environment-based secrets. Automated unit tests, linting and a formal schema are not yet present.')]:
    doc.add_heading(h,2); para(doc,t)
doc.add_heading('16 Testing Strategy',1)
para(doc,'Executed checks included the production build, health/weather/reverse-geocoding requests, Arabic search, browser interaction, mobile layout measurements and provider-failure diagnosis. No automated unit suite, coverage report, load test, formal accessibility audit or penetration test was found.')
tests=[['TC-01','FR-01','Load dashboard','Open root','Weather cards render','Rendered in browser','Passed'],['TC-02','FR-02','Search Arabic “دبي”','Enter text','Result includes country','Returned and rendered','Passed'],['TC-03','FR-04','Use GPS','Permission','Area name appears','Permission-dependent','Partial'],['TC-04','FR-05','Hourly strip','Phone width','Only strip scrolls','Body width equals viewport','Passed'],['TC-05','FR-06','Save place','Result exists','Row remains after refresh','Implemented','Passed'],['TC-06','FR-08','Map marker','Tiles available','Marker follows selection','Implemented','Partial'],['TC-07','FR-09','Mibyan response','Valid project key','Model completion','Current key rejected','Failed prerequisite'],['TC-08','FR-10','Provider failure','Invalid/limited key','Fallback appears','Observed','Passed'],['TC-09','FR-09','Chat scroll','Several messages','Last message visible','Auto-scroll verified','Passed'],['TC-10','NFR','Production build','Dependencies','Build exits cleanly','npm run build passed','Passed']]
table(doc,['ID','Req.','Test','Precondition','Expected','Actual','Status'],tests)
doc.add_heading('16.1 Coverage and Bug Tracking',2); para(doc,'No coverage percentage is claimed. Future coverage should target route validation, search deduplication, localStorage, loading/error states, provider failures and map cleanup. A future bug record should contain ID, reproduction, expected/actual result, severity, fix and regression test.')
doc.add_heading('17 Software Quality Evaluation',1)
table(doc,['Dimension','Assessment','Classification'],[['Correctness','Core weather/search flows are implemented; AI depends on key','Partially verified'],['Readability','Names are clear; App.tsx is dense','Partially verified'],['Maintainability','Server boundary and env config are useful','Partially verified'],['Reliability','Chat fallback prevents total failure','Verified for fallback'],['Usability','No login, disambiguation and mobile containment','Verified for tested view'],['Performance','No benchmark','Not tested'],['Security','Server-side key handling; no formal audit','Partially verified'],['Compatibility','Phone viewport checked; device matrix incomplete','Partially verified']])
doc.add_heading('18 Security and Privacy',1)
para(doc,'Mibyan credentials are read by Express from environment variables and are not sent to the React browser. Coordinate and query inputs are validated before provider calls. CORS is permissive for local development and should be restricted for deployment. Location is used to retrieve weather and reverse geocoding; a future analytics or account layer would require clear privacy choices. No authentication or authorization exists because no accounts exist.')
doc.add_heading('19 Challenges and Limitations',1)
bullets(doc,['No database, accounts or cross-device synchronisation.','Weather values depend on upstream provider location, timing and model.','The map is a base map with marker; weather raster layers and storms are not implemented.','Mibyan adapter exists, but the current project key returned HTTP 401 during verification.','Fallback chat is bounded and is not equivalent to model reasoning or web search.','App.tsx needs feature extraction as scope grows.','No automated tests, coverage, load, accessibility or security audit was found.'])
doc.add_heading('20 Future Improvements',1)
table(doc,['Improvement','Problem','Benefit','Complexity','Priority'],[['Validate/rotate Mibyan key','401 prevents live completions','Real AI answers','Low','High'],['Extract service modules','Dense App.tsx','Easier tests/providers','Medium','High'],['Automated tests','Manual verification only','Regression evidence','Medium','High'],['Weather map layers','Base marker map','Storm/rain context','High','Medium'],['Comparison view','Limited saved-row comparison','Better planning','Medium','Medium'],['Account sync','Device-only storage','Cross-device favourites','High','Low'],['Rate limiting/schema validation','Public routes','Safer deployment','Medium','High'],['Accessibility audit','Partial checks','Better inclusive access','Medium','Medium']])
doc.add_heading('21 Conclusion',1)
para(doc,'Al-Tarqi SkyCast demonstrates a complete small-scale web construction project with a clear visitor workflow: enter the dashboard, identify a place, read current and hourly weather, save places, inspect a map and ask for an interpretation. The solution is grounded in React/TypeScript, Node/Express, Open-Meteo, Leaflet/OpenStreetMap, localStorage and an OpenAI-compatible Mibyan adapter.')
para(doc,'The project demonstrates practical Software Construction through API boundaries, environment configuration, reusable UI functions, responsive containment, error handling and iterative verification. The report also preserves an important engineering distinction: the Mibyan integration is implemented according to its documented contract, but the current key failed authentication, so live model quality is not claimed. The fallback provides continuity but remains limited.')
doc.add_heading('22 References',1)
for ref in ['Mibyan AI. (n.d.). API overview. https://docs.mibyanai.com/api-reference/introduction','Mibyan AI. (n.d.). API quickstart. https://docs.mibyanai.com/api-reference/quickstart','Open-Meteo. (n.d.). Weather API documentation. https://open-meteo.com/en/docs','Open-Meteo. (n.d.). Geocoding API documentation. https://open-meteo.com/en/docs/geocoding-api','Leaflet. (n.d.). Leaflet reference. https://leafletjs.com/reference.html','OpenStreetMap Foundation. (n.d.). Copyright and attribution. https://www.openstreetmap.org/copyright','React. (n.d.). React documentation. https://react.dev/','TypeScript. (n.d.). TypeScript handbook. https://www.typescriptlang.org/docs/handbook/intro.html','Vite. (n.d.). Vite guide. https://vite.dev/guide/','Express. (n.d.). Express documentation. https://expressjs.com/','MDN Web Docs. (n.d.). Window localStorage. https://developer.mozilla.org/en-US/docs/Web/API/Window/localStorage','MDN Web Docs. (n.d.). Geolocation API. https://developer.mozilla.org/en-US/docs/Web/API/Geolocation_API']:
    p=doc.add_paragraph(style='Body Text'); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25); p.add_run(ref)
doc.add_heading('Appendix A Technical Details',1); table(doc,['Item','Detail'],[['Weather','GET /api/weather?latitude=...&longitude=...'],['Search','GET /api/search?q=...'],['Chat','POST /api/chat with question, weather and location JSON'],['Mibyan config','MIBYAN_API_KEY, MIBYAN_BASE_URL, MIBYAN_MODEL'],['Storage','skycast:place, skycast:favorites, skycast:theme'],['Database','None found in current implementation']])
doc.add_heading('Appendix B Evidence Notes',1); para(doc,'This report was generated from the project source files, package configuration, executed build/API checks, browser inspection at a phone-sized viewport and the captured mobile screenshot. Student and course metadata remain editable placeholders because they were not supplied.')
doc.add_heading('Appendix C Diagram Index',1)
for pth in sorted(D.glob('*.png')): para(doc,pth.name)
out=OUT/'Software Construction Project Report.docx'; doc.save(out); print(out)

